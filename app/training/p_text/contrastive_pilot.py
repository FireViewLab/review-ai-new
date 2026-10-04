"""Validation utilities for the contrastive hard-example annotation pilot.

The JSONL annotation file is the source of truth.  This module never mutates
the pilot or any forbidden corpus; the builder is a separate module.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterable, Mapping, Sequence

from app.integrations.similarity import bigram_dice, canonical_text, sparse_cosine, tfidf_vectors


CATEGORIES = (
    "EXPERIENCE_FREE_CLAIM",
    "UNUSED_RECOMMENDATION",
    "NATURAL_PROMOTIONAL",
    "PROMOTION_EXPERIENCE_MIX",
    "FEATURE_AS_PERSONAL_EXPERIENCE",
)
SUPPORTED_CATEGORIES = frozenset((*CATEGORIES, "EXAGGERATED_PRAISE_WITHOUT_EVIDENCE"))
SUPPORTED_DOMAINS = frozenset(
    {
        "cosmetics", "health_supplement", "pharma_otc", "food", "fashion",
        "household", "electronics",
        # Legacy pilot values remain valid during migration.
        "cosmetics_health", "mixed",
    }
)
REQUIRED_COLUMNS = (
    "sample_id", "content", "label", "category", "secondary_categories",
    "domain", "source_type", "parent_pair_id", "generation_family_id",
    "difficulty", "review_status", "notes", "author_id", "reviewer_ids",
    "split", "seed", "created_at", "source_reference", "creation_method",
    "original_source_type", "is_llm_generated", "pair_relation",
    "length_bucket", "style",
)
EXTENDED_METADATA_COLUMNS = (
    "praise_intensity", "evidence_level", "human_reviewed", "human_approved",
)
EXPORT_COLUMNS = (*REQUIRED_COLUMNS, *EXTENDED_METADATA_COLUMNS)
PRAISE_INTENSITIES = frozenset({"low", "medium", "extreme"})
EVIDENCE_LEVELS = frozenset({"none", "limited", "sufficient"})
HEALTH_KEYWORDS = (
    "비타민", "영양제", "유산균", "오메가3", "홍삼", "효과", "피로",
    "건강", "성분", "약", "복용", "추천", "최고", "완벽", "만족",
)
LABELS = frozenset({"NORMAL", "SUSPICIOUS"})
SOURCE_TYPES = frozenset(
    {"human_written", "human_curated", "synthetic", "augmented", "PENDING_HUMAN"}
)
REVIEW_STATUSES = frozenset({"DRAFT", "AGREED", "ADJUDICATED", "REJECTED"})
TRAINING_READY_STATUSES = frozenset({"AGREED", "ADJUDICATED"})
MAX_GENERATION_FAMILY_PAIRS = 4
NEAR_TEMPLATE_THRESHOLD = 0.85
CROSS_CORPUS_REPORT_FLOOR = 0.80


@dataclass(frozen=True, slots=True)
class ForbiddenRecord:
    source: str
    row_identifier: str
    content: str
    corpus_kind: str


@dataclass(frozen=True, slots=True)
class CrossCorpusMatch:
    pilot_sample_id: str
    forbidden_source: str
    forbidden_row_identifier: str
    corpus_kind: str
    similarity: float
    band: str
    candidate: bool


def normalized_text(value: str) -> str:
    """Normalize for duplicate checks without changing stored content."""

    value = unicodedata.normalize("NFKC", value).casefold().strip()
    return re.sub(r"\s+", " ", value)


@dataclass
class ValidationReport:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    counts: dict[str, Any] = field(default_factory=dict)
    cross_corpus_matches: list[CrossCorpusMatch] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return not self.errors


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def validate_pilot(
    rows: Sequence[Mapping[str, Any]],
    *,
    forbidden_contents: Iterable[str] = (),
    forbidden_records: Sequence[ForbiddenRecord] = (),
    expected_pairs_per_category: int | None = 10,
    max_generation_family_pairs: int = MAX_GENERATION_FAMILY_PAIRS,
) -> ValidationReport:
    report = ValidationReport()
    required = set(REQUIRED_COLUMNS)
    for index, row in enumerate(rows, start=1):
        missing = sorted(required.difference(row))
        if missing:
            report.errors.append(f"row {index}: missing required columns: {', '.join(missing)}")

    ids = [str(row.get("sample_id", "")) for row in rows]
    duplicate_ids = sorted(key for key, count in Counter(ids).items() if key and count > 1)
    if duplicate_ids:
        report.errors.append("duplicate sample_id: " + ", ".join(duplicate_ids))

    pairs: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for row in rows:
        pairs[str(row.get("parent_pair_id", ""))].append(row)
    for pair_id, pair_rows in pairs.items():
        if not pair_id:
            report.errors.append("blank parent_pair_id")
            continue
        if len(pair_rows) != 2:
            report.errors.append(f"{pair_id}: expected exactly 2 rows, got {len(pair_rows)}")
        labels = Counter(str(row.get("label", "")) for row in pair_rows)
        if labels != Counter({"NORMAL": 1, "SUSPICIOUS": 1}):
            report.errors.append(f"{pair_id}: pair must contain one NORMAL and one SUSPICIOUS")
        categories = {str(row.get("category", "")) for row in pair_rows}
        if len(categories) != 1:
            report.errors.append(f"{pair_id}: pair category mismatch")
        source_types = {str(row.get("source_type", "")) for row in pair_rows}
        if len(source_types) != 1:
            report.errors.append(f"{pair_id}: pair source_type mismatch")

    category_pairs = Counter(
        str(pair_rows[0].get("category", ""))
        for pair_rows in pairs.values()
        if pair_rows
    )
    if expected_pairs_per_category is not None:
        if len(rows) != 100:
            report.errors.append(f"expected 100 rows, got {len(rows)}")
        if len(pairs) != 50:
            report.errors.append(f"expected 50 pairs, got {len(pairs)}")
        unknown_categories = sorted(set(category_pairs).difference(SUPPORTED_CATEGORIES))
        if unknown_categories:
            report.errors.append("unknown categories: " + ", ".join(unknown_categories))
        for category in CATEGORIES:
            if category_pairs[category] != expected_pairs_per_category:
                report.errors.append(
                    f"{category}: expected {expected_pairs_per_category} pairs, "
                    f"got {category_pairs[category]}"
                )

    source_rows = Counter(str(row.get("source_type", "")) for row in rows)
    for index, row in enumerate(rows, start=1):
        source_type = str(row.get("source_type", ""))
        status = str(row.get("review_status", ""))
        content = str(row.get("content", ""))
        split = str(row.get("split", ""))
        category = str(row.get("category", ""))
        domain = str(row.get("domain", ""))
        praise_intensity = row.get("praise_intensity", "")
        evidence_level = row.get("evidence_level", "")
        is_llm_generated = row.get("is_llm_generated", False)
        human_reviewed = row.get("human_reviewed", False)
        human_approved = row.get("human_approved", False)
        if str(row.get("label", "")) not in LABELS:
            report.errors.append(f"row {index}: invalid label")
        if category not in SUPPORTED_CATEGORIES:
            report.errors.append(f"row {index}: invalid category {category!r}")
        if domain not in SUPPORTED_DOMAINS:
            report.errors.append(f"row {index}: invalid domain {domain!r}")
        if source_type not in SOURCE_TYPES:
            report.errors.append(f"row {index}: invalid source_type {source_type!r}")
        if status not in REVIEW_STATUSES:
            report.errors.append(f"row {index}: invalid review_status {status!r}")
        if praise_intensity not in ("", None) and praise_intensity not in PRAISE_INTENSITIES:
            report.errors.append(f"row {index}: invalid praise_intensity {praise_intensity!r}")
        if evidence_level not in ("", None) and evidence_level not in EVIDENCE_LEVELS:
            report.errors.append(f"row {index}: invalid evidence_level {evidence_level!r}")
        for name, value in (
            ("is_llm_generated", is_llm_generated),
            ("human_reviewed", human_reviewed),
            ("human_approved", human_approved),
        ):
            if not isinstance(value, bool):
                report.errors.append(f"row {index}: {name} must be boolean")
        if isinstance(is_llm_generated, bool):
            if source_type == "synthetic" and not is_llm_generated:
                report.errors.append(f"row {index}: synthetic row must retain is_llm_generated=true")
            if source_type in {"human_written", "human_curated", "PENDING_HUMAN"} and is_llm_generated:
                report.errors.append(
                    f"row {index}: is_llm_generated=true conflicts with source_type={source_type}"
                )
        if isinstance(human_approved, bool) and isinstance(human_reviewed, bool):
            if human_approved and not human_reviewed:
                report.errors.append(f"row {index}: human_approved=true requires human_reviewed=true")
        if category == "EXAGGERATED_PRAISE_WITHOUT_EVIDENCE":
            if praise_intensity in ("", None) or evidence_level in ("", None):
                report.errors.append(
                    f"row {index}: exaggerated-praise rows require praise_intensity and evidence_level"
                )
            elif str(row.get("label", "")) == "SUSPICIOUS" and not (
                praise_intensity == "extreme" and evidence_level in {"none", "limited"}
            ):
                report.errors.append(
                    f"row {index}: suspicious exaggerated praise requires extreme praise "
                    "with none/limited evidence"
                )
        if not content.strip():
            if source_type != "PENDING_HUMAN":
                report.errors.append(f"row {index}: blank content must be PENDING_HUMAN")
            if status in TRAINING_READY_STATUSES or split.strip():
                report.errors.append(f"row {index}: blank placeholder marked training-ready")
        elif source_type == "PENDING_HUMAN":
            report.errors.append(f"row {index}: populated content cannot remain PENDING_HUMAN")

    synthetic_pair_count = sum(
        1
        for pair_rows in pairs.values()
        if pair_rows and {str(row.get("source_type", "")) for row in pair_rows}
        <= {"synthetic", "augmented"}
    )
    pending_pair_count = sum(
        1
        for pair_rows in pairs.values()
        if pair_rows and {str(row.get("source_type", "")) for row in pair_rows}
        == {"PENDING_HUMAN"}
    )
    human_pair_count = sum(
        1
        for pair_rows in pairs.values()
        if pair_rows and {str(row.get("source_type", "")) for row in pair_rows}
        <= {"human_written", "human_curated"}
    )
    if expected_pairs_per_category is not None:
        if synthetic_pair_count > 20:
            report.errors.append(f"synthetic/augmented pair limit exceeded: {synthetic_pair_count}")
        if human_pair_count + pending_pair_count < 30:
            report.errors.append("human or pending-human pair target is below 30")
        for category in CATEGORIES:
            category_pair_rows = [
                pair_rows for pair_rows in pairs.values()
                if pair_rows and str(pair_rows[0].get("category", "")) == category
            ]
            synthetic_or_augmented = sum(
                {str(row.get("source_type", "")) for row in pair_rows}
                <= {"synthetic", "augmented"}
                for pair_rows in category_pair_rows
            )
            human_or_pending = sum(
                {str(row.get("source_type", "")) for row in pair_rows}
                <= {"human_written", "human_curated", "PENDING_HUMAN"}
                for pair_rows in category_pair_rows
            )
            if synthetic_or_augmented != 4:
                report.errors.append(
                    f"{category}: expected 4 synthetic/augmented pairs, got {synthetic_or_augmented}"
                )
            if human_or_pending != 6:
                report.errors.append(
                    f"{category}: expected 6 human/pending pairs, got {human_or_pending}"
                )

    populated = [row for row in rows if str(row.get("content", "")).strip()]
    exact_groups = _duplicate_groups(populated, lambda text: text)
    normalized_groups = _duplicate_groups(populated, normalized_text)
    conflicts = _label_conflicts(populated)
    if exact_groups:
        report.errors.append(f"exact duplicate groups: {len(exact_groups)}")
    if normalized_groups:
        report.errors.append(f"normalized duplicate groups: {len(normalized_groups)}")
    if conflicts:
        report.errors.append(f"label conflict groups: {len(conflicts)}")

    keyword_only_pairs = _keyword_only_pair_candidates(pairs)
    if keyword_only_pairs:
        report.errors.append("keyword-only minimal pair candidates: " + ", ".join(keyword_only_pairs))

    family_pairs: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        family_pairs[str(row.get("generation_family_id", ""))].add(
            str(row.get("parent_pair_id", ""))
        )
    over_limit_families = {
        family: len(pair_ids)
        for family, pair_ids in family_pairs.items()
        if family and len(pair_ids) > max_generation_family_pairs
    }
    if over_limit_families:
        report.errors.append(f"generation family limit exceeded: {over_limit_families}")

    domain_rows = Counter(str(row.get("domain", "")) for row in populated)
    if (expected_pairs_per_category is not None and populated
            and max(domain_rows.values(), default=0) / len(populated) > 0.30):
        report.errors.append("domain concentration exceeds 30% among populated rows")

    supplied_forbidden = list(forbidden_contents)
    supplied_forbidden.extend(record.content for record in forbidden_records)
    forbidden_exact = {text for text in supplied_forbidden if text}
    forbidden_normalized = {normalized_text(text) for text in forbidden_exact}
    exact_overlap = [str(row["sample_id"]) for row in populated if str(row["content"]) in forbidden_exact]
    normalized_overlap = [
        str(row["sample_id"])
        for row in populated
        if normalized_text(str(row["content"])) in forbidden_normalized
    ]
    if exact_overlap:
        report.errors.append("forbidden exact overlap: " + ", ".join(exact_overlap))
    if normalized_overlap:
        report.errors.append("forbidden normalized overlap: " + ", ".join(normalized_overlap))

    cross_corpus_matches = cross_corpus_near_template_matches(populated, forbidden_records)
    report.cross_corpus_matches = cross_corpus_matches
    protected_matches = [
        match for match in cross_corpus_matches
        if match.candidate and match.corpus_kind in {"blind", "challenge"}
    ]
    training_matches = [
        match for match in cross_corpus_matches
        if match.candidate and match.corpus_kind == "training"
    ]
    if protected_matches:
        report.errors.append(
            f"blind/challenge cross-corpus similarity >= {NEAR_TEMPLATE_THRESHOLD}: "
            f"{len(protected_matches)}"
        )
    if training_matches:
        report.warnings.append(
            f"training cross-corpus similarity >= {NEAR_TEMPLATE_THRESHOLD}: "
            f"{len(training_matches)}"
        )

    near_template_pairs = _cross_family_near_templates(populated)
    if near_template_pairs:
        report.warnings.append(
            f"cross-family near-template candidates >= {NEAR_TEMPLATE_THRESHOLD}: "
            f"{len(near_template_pairs)}"
        )

    keyword_balance = health_keyword_balance(populated)
    report.counts = {
        "rows": len(rows),
        "pairs": len(pairs),
        "category_pairs": dict(category_pairs),
        "source_type_rows": dict(source_rows),
        "synthetic_or_augmented_pairs": synthetic_pair_count,
        "human_pairs": human_pair_count,
        "pending_human_pairs": pending_pair_count,
        "domain_rows_populated": dict(domain_rows),
        "difficulty_rows": dict(Counter(str(row.get("difficulty", "")) for row in rows)),
        "exact_duplicate_groups": len(exact_groups),
        "normalized_duplicate_groups": len(normalized_groups),
        "label_conflict_groups": len(conflicts),
        "keyword_only_pair_candidates": len(keyword_only_pairs),
        "forbidden_exact_overlap": len(exact_overlap),
        "forbidden_normalized_overlap": len(normalized_overlap),
        "cross_corpus_similarity_bands": dict(Counter(match.band for match in cross_corpus_matches)),
        "cross_corpus_ge_085_by_kind": dict(
            Counter(match.corpus_kind for match in cross_corpus_matches if match.candidate)
        ),
        "cross_family_near_template_candidates": len(near_template_pairs),
        "training_ready_rows": sum(
            str(row.get("review_status", "")) in TRAINING_READY_STATUSES
            and bool(str(row.get("content", "")).strip())
            for row in rows
        ),
        "health_keyword_balance": keyword_balance,
    }
    return report


def health_keyword_balance(rows: Sequence[Mapping[str, Any]]) -> dict[str, dict[str, int]]:
    """Report label-wise keyword counts without using keywords to assign labels."""

    balance: dict[str, dict[str, int]] = {}
    for keyword in HEALTH_KEYWORDS:
        counts = {"NORMAL": 0, "SUSPICIOUS": 0}
        for row in rows:
            label = str(row.get("label", ""))
            if label in counts and _contains_health_keyword(str(row.get("content", "")), keyword):
                counts[label] += 1
        balance[keyword] = counts
    return balance


def _contains_health_keyword(content: str, keyword: str) -> bool:
    if keyword != "약":
        return keyword in content
    return bool(
        re.search(
            r"(?<![0-9A-Za-z가-힣])약(?:을|를|이|가|은|는|과|와|의|으로|에|도|만)?"
            r"(?![0-9A-Za-z가-힣])",
            content,
        )
    )


def read_forbidden_contents(paths: Iterable[Path]) -> list[str]:
    """Read text fields from existing corpora without modifying them."""

    contents: list[str] = []
    for path in paths:
        suffix = path.suffix.casefold()
        if suffix == ".jsonl":
            for row in load_jsonl(path):
                contents.extend(_content_values(row))
        elif suffix == ".json":
            with path.open("r", encoding="utf-8") as handle:
                contents.extend(_content_values(json.load(handle)))
        elif suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                for row in csv.DictReader(handle):
                    contents.extend(_content_values(row))
        elif suffix == ".xlsx":
            contents.extend(_read_xlsx_contents(path))
    return contents


def read_forbidden_records(paths: Iterable[Path]) -> list[ForbiddenRecord]:
    """Read auditable content records and preserve file and row identity."""

    records: list[ForbiddenRecord] = []
    for path in paths:
        kind = _corpus_kind(path)
        suffix = path.suffix.casefold()
        if suffix == ".xlsx":
            records.extend(_read_xlsx_records(path, kind))
        elif suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                for index, row in enumerate(csv.DictReader(handle), start=2):
                    records.extend(_mapping_records(path, kind, row, f"row:{index}"))
        elif suffix == ".jsonl":
            for index, row in enumerate(load_jsonl(path), start=1):
                if isinstance(row, Mapping):
                    records.extend(_mapping_records(path, kind, row, f"line:{index}"))
        elif suffix == ".json":
            with path.open("r", encoding="utf-8") as handle:
                value = json.load(handle)
            mappings = value if isinstance(value, list) else [value]
            for index, row in enumerate(mappings, start=1):
                if isinstance(row, Mapping):
                    records.extend(_mapping_records(path, kind, row, f"item:{index}"))
    return records


def cross_corpus_near_template_matches(
    pilot_rows: Sequence[Mapping[str, Any]],
    forbidden_records: Sequence[ForbiddenRecord],
    *,
    report_floor: float = CROSS_CORPUS_REPORT_FLOOR,
) -> list[CrossCorpusMatch]:
    """Return raw Hybrid-A similarities in fixed reporting bands."""

    pilot = [
        row for row in pilot_rows
        if str(row.get("content", "")).strip()
        and str(row.get("source_type", "")) in {"synthetic", "augmented"}
    ]
    forbidden = [record for record in forbidden_records if record.content.strip()]
    if not pilot or not forbidden:
        return []
    texts = [str(row["content"]) for row in pilot] + [record.content for record in forbidden]
    canonical = [canonical_text(text) for text in texts]
    vectors = tfidf_vectors(canonical)
    matches: list[CrossCorpusMatch] = []
    offset = len(pilot)
    for pilot_index, row in enumerate(pilot):
        for forbidden_index, record in enumerate(forbidden):
            right_index = offset + forbidden_index
            if canonical[pilot_index] == canonical[right_index]:
                similarity = 1.0
            else:
                cosine = sparse_cosine(vectors[pilot_index], vectors[right_index])
                dice = bigram_dice(canonical[pilot_index], canonical[right_index])
                if cosine is None or dice is None:
                    continue
                similarity = 0.40 * cosine + 0.60 * dice
            if similarity < report_floor:
                continue
            if similarity >= 0.90:
                band = ">=0.90"
            elif similarity >= 0.85:
                band = "0.85-0.90"
            else:
                band = "0.80-0.85"
            matches.append(
                CrossCorpusMatch(
                    pilot_sample_id=str(row["sample_id"]),
                    forbidden_source=record.source,
                    forbidden_row_identifier=record.row_identifier,
                    corpus_kind=record.corpus_kind,
                    similarity=similarity,
                    band=band,
                    candidate=similarity >= NEAR_TEMPLATE_THRESHOLD,
                )
            )
    return sorted(
        matches,
        key=lambda match: (-match.similarity, match.pilot_sample_id,
                           match.forbidden_source, match.forbidden_row_identifier),
    )


def discover_forbidden_paths(
    roots: Iterable[Path], *, pilot_root: Path | None = None
) -> list[Path]:
    """Discover supported corpora while explicitly excluding the pilot tree."""

    resolved_pilot = pilot_root.resolve() if pilot_root is not None else None
    found: list[Path] = []
    for root in roots:
        for path in root.rglob("*"):
            if not path.is_file() or path.suffix.casefold() not in {".xlsx", ".csv", ".jsonl", ".json"}:
                continue
            resolved = path.resolve()
            if resolved_pilot is not None and (resolved == resolved_pilot or resolved_pilot in resolved.parents):
                continue
            found.append(path)
    return sorted(found)


def _content_values(value: Any) -> list[str]:
    found: list[str] = []
    if isinstance(value, Mapping):
        for key, child in value.items():
            if str(key).casefold() in {"content", "review", "text"} and isinstance(child, str):
                if child.strip():
                    found.append(child)
            elif isinstance(child, (Mapping, list)):
                found.extend(_content_values(child))
    elif isinstance(value, list):
        for child in value:
            found.extend(_content_values(child))
    return found


def _read_xlsx_contents(path: Path) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to scan forbidden .xlsx corpora") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    found: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip().casefold() if value is not None else "" for value in next(rows, ())]
            indexes = [index for index, name in enumerate(headers) if name in {"content", "review", "text"}]
            for values in rows:
                for index in indexes:
                    if index < len(values) and isinstance(values[index], str) and values[index].strip():
                        found.append(values[index])
    finally:
        workbook.close()
    return found


def _read_xlsx_records(path: Path, kind: str) -> list[ForbiddenRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to scan forbidden .xlsx corpora") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[ForbiddenRecord] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            raw_headers = next(rows, ())
            headers = [str(value).strip() if value is not None else "" for value in raw_headers]
            lowered = [header.casefold() for header in headers]
            for excel_row, values in enumerate(rows, start=2):
                row = dict(zip(headers, values))
                fallback = f"{sheet.title}!row:{excel_row}"
                for content_name in ("content", "review", "text"):
                    if content_name not in lowered:
                        continue
                    value = values[lowered.index(content_name)]
                    if isinstance(value, str) and value.strip():
                        identifier = _row_identifier(row, fallback)
                        records.append(ForbiddenRecord(str(path), identifier, value, kind))
    finally:
        workbook.close()
    return records


def _mapping_records(
    path: Path, kind: str, row: Mapping[str, Any], fallback: str
) -> list[ForbiddenRecord]:
    lowered = {str(key).casefold(): value for key, value in row.items()}
    records: list[ForbiddenRecord] = []
    for name in ("content", "review", "text"):
        value = lowered.get(name)
        if isinstance(value, str) and value.strip():
            records.append(ForbiddenRecord(str(path), _row_identifier(row, fallback), value, kind))
    return records


def _row_identifier(row: Mapping[str, Any], fallback: str) -> str:
    lowered = {str(key).casefold(): value for key, value in row.items()}
    for name in ("blind_id", "sample_id", "review_id", "source_id", "source_excel_row", "master_id"):
        value = lowered.get(name)
        if value is not None and str(value).strip():
            return f"{name}:{value}"
    return fallback


def _corpus_kind(path: Path) -> str:
    name = path.name.casefold()
    if "blind" in name:
        return "blind"
    if "challenge" in name:
        return "challenge"
    return "training"


def _duplicate_groups(rows: Sequence[Mapping[str, Any]], key_function) -> dict[str, list[str]]:
    groups: dict[str, list[str]] = defaultdict(list)
    for row in rows:
        groups[key_function(str(row["content"]))].append(str(row["sample_id"]))
    return {key: ids for key, ids in groups.items() if key and len(ids) > 1}


def _label_conflicts(rows: Sequence[Mapping[str, Any]]) -> dict[str, set[str]]:
    groups: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        groups[normalized_text(str(row["content"]))].add(str(row["label"]))
    return {key: labels for key, labels in groups.items() if len(labels) > 1}


def _keyword_only_pair_candidates(pairs: Mapping[str, Sequence[Mapping[str, Any]]]) -> list[str]:
    candidates: list[str] = []
    token_pattern = re.compile(r"[0-9A-Za-z가-힣]+")
    for pair_id, pair_rows in pairs.items():
        if len(pair_rows) != 2 or any(not str(row.get("content", "")).strip() for row in pair_rows):
            continue
        left, right = (str(row["content"]) for row in pair_rows)
        left_tokens = token_pattern.findall(normalized_text(left))
        right_tokens = token_pattern.findall(normalized_text(right))
        changed = Counter(left_tokens) - Counter(right_tokens)
        changed.update(Counter(right_tokens) - Counter(left_tokens))
        dice = bigram_dice(canonical_text(left), canonical_text(right)) or 0.0
        sequence_similarity = SequenceMatcher(
            None, canonical_text(left), canonical_text(right), autojunk=False
        ).ratio()
        if sum(changed.values()) <= 2 and (dice >= 0.90 or sequence_similarity >= 0.75):
            candidates.append(pair_id)
    return sorted(candidates)


def _cross_family_near_templates(rows: Sequence[Mapping[str, Any]]) -> list[tuple[str, str, float]]:
    if len(rows) < 2:
        return []
    canonical = [canonical_text(str(row["content"])) for row in rows]
    vectors = tfidf_vectors(canonical)
    candidates: list[tuple[str, str, float]] = []
    for left in range(len(rows)):
        for right in range(left + 1, len(rows)):
            if rows[left].get("generation_family_id") == rows[right].get("generation_family_id"):
                continue
            cosine = sparse_cosine(vectors[left], vectors[right])
            dice = bigram_dice(canonical[left], canonical[right])
            if cosine is None or dice is None:
                continue
            similarity = 0.40 * cosine + 0.60 * dice
            if similarity >= NEAR_TEMPLATE_THRESHOLD:
                candidates.append((str(rows[left]["sample_id"]), str(rows[right]["sample_id"]), similarity))
    return candidates


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pilot_path", type=Path)
    parser.add_argument("--forbidden", type=Path, action="append", default=[])
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    rows = load_jsonl(args.pilot_path)
    forbidden_records = read_forbidden_records(args.forbidden)
    report = validate_pilot(rows, forbidden_records=forbidden_records)
    print(json.dumps({"passed": report.passed, "errors": report.errors,
                      "warnings": report.warnings, "counts": report.counts,
                      "cross_corpus_matches": [
                          {
                              "pilot_sample_id": match.pilot_sample_id,
                              "forbidden_source": match.forbidden_source,
                              "forbidden_row_identifier": match.forbidden_row_identifier,
                              "corpus_kind": match.corpus_kind,
                              "similarity": match.similarity,
                              "band": match.band,
                              "candidate": match.candidate,
                          }
                          for match in report.cross_corpus_matches
                      ]},
                     ensure_ascii=False, indent=2))
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
