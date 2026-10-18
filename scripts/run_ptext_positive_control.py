"""Run a bounded P_text runtime control and human-reviewed exploratory panel."""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
from typing import Callable, Iterable, Mapping

from app.analyzers.p_text import DEFAULT_MODEL_PATH, predict_text_score
from app.training.p_text.data import normalize_text, prepare_records, read_review_master


ROOT = Path(__file__).resolve().parents[1]
PREDICTIONS_PATH = ROOT / "reports/ptext-koelectra-v1-2epoch-20260929/predictions.jsonl"
V1_SOURCE_WORKBOOK = Path(
    r"C:\Users\kimha\Projects\ai-ptext-data-prep\outputs\ptext_training_dataset_v2.xlsx"
)
V3_WORKBOOK = next((ROOT / "data").glob("ReView_Integrated_Review_Dataset_v3_2*.xlsx"))
V4_WORKBOOK = next((ROOT / "data").glob("ReView_Integrated_Review_Dataset_v4_3*.xlsx"))
DEFAULT_JSON = ROOT / "reports/ptext_positive_control_20261018.json"
DEFAULT_MARKDOWN = ROOT / "reports/ptext_positive_control_20261018.md"
SYNTHETIC_TYPES = (
    "COPY_VARIATION",
    "EXPERIENCE_FREE_CLAIM",
    "NO_USE_ASSERTIVE_RECOMMENDATION",
    "PRODUCT_DESCRIPTION_REPLACEMENT",
    "PROMOTIONAL_TEMPLATE",
)


def stable_digest(*parts: object) -> str:
    return hashlib.sha256("\x1f".join(str(part) for part in parts).encode("utf-8")).hexdigest()


def stable_select(rows: Iterable[Mapping[str, object]], count: int) -> list[Mapping[str, object]]:
    """Select the smallest stable identities independently of input order."""
    if count < 0:
        raise ValueError("count must not be negative")
    ranked = sorted(
        rows,
        key=lambda row: stable_digest(
            row.get("evaluation_set", ""),
            row.get("synthetic_type", ""),
            row.get("platform", ""),
            row.get("product_id", ""),
            row.get("review_id", ""),
            row.get("excel_row", ""),
            row.get("text", ""),
        ),
    )
    return ranked[:count]


def load_predictions(path: Path) -> list[dict[str, object]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def select_panel_a(predictions: Iterable[Mapping[str, object]]) -> list[dict[str, object]]:
    rows = list(predictions)
    normal_candidates = [
        row for row in rows
        if row.get("evaluation_set") == "real_NORMAL_human"
        and row.get("actual_label") == 0
        and row.get("predicted_default") == 0
    ]
    selected: list[dict[str, object]] = []
    for index, row in enumerate(stable_select(normal_candidates, 5), start=1):
        selected.append({
            **row,
            "control_id": f"A-NORMAL-{index:02d}-{stable_digest(row.get('review_id'), row.get('excel_row'))[:8]}",
            "source_type": "previous_model_real_normal",
            "provenance": "previous-model expected; source label real_NORMAL_human",
            "expected_label": "NORMAL",
        })
    for index, synthetic_type in enumerate(SYNTHETIC_TYPES, start=1):
        candidates = [
            row for row in rows
            if row.get("synthetic_type") == synthetic_type
            and row.get("actual_label") == 1
            and row.get("predicted_default") == 1
        ]
        chosen = stable_select(candidates, 1)
        if len(chosen) != 1:
            raise RuntimeError(f"could not select one historical positive for {synthetic_type}")
        row = chosen[0]
        selected.append({
            **row,
            "control_id": f"A-SUSPICIOUS-{index:02d}-{stable_digest(row.get('review_id'), row.get('excel_row'))[:8]}",
            "source_type": "previous_model_synthetic_suspicious",
            "provenance": "previous-model expected; synthetic label (not human ground truth)",
            "expected_label": "SUSPICIOUS",
        })
    if len(selected) != 10:
        raise RuntimeError(f"Panel A must contain exactly 10 controls, got {len(selected)}")
    return selected


def resolve_panel_a(selected: list[dict[str, object]], workbook_path: Path) -> list[dict[str, object]]:
    from openpyxl import load_workbook

    needed = {int(row["excel_row"]): row for row in selected}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["dataset"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        resolved: dict[int, dict[str, object]] = {}
        for excel_row, values in enumerate(rows, start=2):
            if excel_row in needed:
                resolved[excel_row] = dict(zip(headers, values, strict=True))
    finally:
        workbook.close()

    output = []
    for control in selected:
        excel_row = int(control["excel_row"])
        source = resolved.get(excel_row)
        if source is None:
            raise RuntimeError(f"source row not found: {excel_row}")
        if (
            str(source.get("review_id")) != str(control.get("review_id"))
            or source.get("label") != control.get("label")
            or not isinstance(source.get("content"), str)
            or not source["content"].strip()
        ):
            raise RuntimeError(f"source identity mismatch at Excel row {excel_row}")
        output.append({**control, "content": source["content"]})
    return output


def select_panel_b(v3_path: Path, v4_path: Path) -> list[dict[str, object]]:
    v3 = prepare_records(read_review_master(str(v3_path)))
    v4 = prepare_records(read_review_master(str(v4_path)))
    v3_suspicious = {normalize_text(row["text"]) for row in v3.examples if row["label"] == 1}
    v4_suspicious = {
        normalize_text(row["text"]): row for row in v4.examples if row["label"] == 1
    }
    inherited = sorted(v3_suspicious.intersection(v4_suspicious))
    if len(v3_suspicious) != 4 or len(inherited) != 4:
        raise RuntimeError(
            "Expected the four v3.2 reviewed SUSPICIOUS texts to map exactly into v4.3"
        )
    normal_rows = [
        {"text": row["text"], "label": row["label"]}
        for row in v4.examples if row["label"] == 0
    ]
    selected = []
    for index, row in enumerate(stable_select(normal_rows, 5), start=1):
        selected.append({
            "panel_id": f"B-NORMAL-{index:02d}-{stable_digest(row['text'])[:8]}",
            "human_label": "NORMAL",
            "content": row["text"],
            "source_type": "human_reviewed_v4_3",
            "provenance": "human-reviewed label; not externally verified authenticity ground truth",
        })
    for index, normalized in enumerate(inherited, start=1):
        row = v4_suspicious[normalized]
        selected.append({
            "panel_id": f"B-SUSPICIOUS-{index:02d}-{stable_digest(row['text'])[:8]}",
            "human_label": "SUSPICIOUS",
            "content": row["text"],
            "source_type": "human_reviewed_v3_2_carried_into_v4_3",
            "provenance": (
                "one of four v3.2 human-reviewed SUSPICIOUS texts, exactly matched in v4.3; "
                "not externally verified fake-review ground truth"
            ),
        })
    if len(selected) != 9:
        raise RuntimeError(f"Panel B must contain exactly 9 rows, got {len(selected)}")
    return selected


def infer_panel_a(
    controls: Iterable[Mapping[str, object]],
    predictor: Callable[[str], Mapping[str, object]] = predict_text_score,
) -> list[dict[str, object]]:
    results = []
    for control in controls:
        prediction = dict(predictor(str(control["content"])))
        current_label = prediction.get("predicted_label")
        historical_probability = float(control["suspicious_probability"])
        results.append({
            "control_id": control["control_id"],
            "source_type": control["source_type"],
            "provenance": control["provenance"],
            "synthetic_type": control.get("synthetic_type"),
            "content": control["content"],
            "character_length": len(str(control["content"])),
            "expected_label": control["expected_label"],
            "historical_probability": historical_probability,
            "current_predicted_label": current_label,
            "current_suspicious_probability": prediction.get("suspicious_probability"),
            "current_text_score": prediction.get("text_score"),
            "label_match": current_label == control["expected_label"],
            "absolute_probability_delta": (
                abs(float(prediction["suspicious_probability"]) - historical_probability)
                if prediction.get("suspicious_probability") is not None else None
            ),
        })
    return results


def infer_panel_b(
    controls: Iterable[Mapping[str, object]],
    predictor: Callable[[str], Mapping[str, object]] = predict_text_score,
) -> list[dict[str, object]]:
    results = []
    for control in controls:
        prediction = dict(predictor(str(control["content"])))
        results.append({
            **control,
            "character_length": len(str(control["content"])),
            "current_predicted_label": prediction.get("predicted_label"),
            "current_suspicious_probability": prediction.get("suspicious_probability"),
            "current_text_score": prediction.get("text_score"),
            "human_model_label_match": prediction.get("predicted_label") == control["human_label"],
        })
    return results


def summarize_panel_a(rows: Iterable[Mapping[str, object]]) -> dict[str, object]:
    records = list(rows)
    expected = Counter(str(row["expected_label"]) for row in records)
    reproduced = Counter(
        str(row["expected_label"]) for row in records if row.get("label_match") is True
    )
    matches = sum(row.get("label_match") is True for row in records)
    return {
        "status": "PASS" if len(records) == 10 and matches == 10 else "FAIL",
        "sample_count": len(records),
        "label_match_count": matches,
        "expected_counts": dict(expected),
        "reproduced_counts": dict(reproduced),
    }


def summarize_panel_b(rows: Iterable[Mapping[str, object]]) -> dict[str, object]:
    records = list(rows)
    result = {}
    for human_label in ("NORMAL", "SUSPICIOUS"):
        group = [row for row in records if row["human_label"] == human_label]
        predictions = Counter(str(row["current_predicted_label"]) for row in group)
        result[human_label] = {
            "sample_count": len(group),
            "model_NORMAL_count": predictions.get("NORMAL", 0),
            "model_SUSPICIOUS_count": predictions.get("SUSPICIOUS", 0),
            "model_unavailable_count": predictions.get("None", 0),
        }
    return result


def _safe(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def _probability(value: object) -> str:
    return "N/A" if value is None else f"{float(value):.9f}"


def render_markdown(report: Mapping[str, object]) -> str:
    a = report["panel_a"]
    b = report["panel_b"]
    lines = [
        "# P_text Positive Control + Human-reviewed Exploratory Panel",
        "",
        "## Purpose",
        "",
        "Panel A checks runtime/model/config reproducibility against previous v1 predictions. "
        "Panel B descriptively observes behavior on a small human-reviewed panel. The two provenance types are not interchangeable.",
        "",
        "## Current runtime/model",
        "",
        f"- Model: `{report['metadata']['model_version']}`",
        "- Runtime: `app.analyzers.p_text.predict_text_score`",
        f"- Inference config: `{json.dumps(report['metadata']['inference_config'], ensure_ascii=False)}`",
        "- Current `text_score` applies T=2; predicted label uses raw probability and threshold 0.5.",
        "",
        "## Panel A source/provenance",
        "",
        "- NORMAL: previous v1 `real_NORMAL_human` rows with historical `predicted_default=0`.",
        "- SUSPICIOUS: one historical `predicted_default=1` synthetic row per synthetic type.",
        "- Content was resolved by exact Excel row, review ID, and label checks against the v1 source workbook.",
        "- Synthetic SUSPICIOUS is not human or external ground truth.",
        "",
        "## Panel A results",
        "",
        f"**{a['summary']['status']} — {a['summary']['label_match_count']}/{a['summary']['sample_count']} labels reproduced.**",
        "",
        "| Control | Type | Expected | Current | Current p(susp.) | Text score | |Δp| | Match |",
        "|---|---|---|---|---:|---:|---:|---|",
    ]
    for row in a["samples"]:
        lines.append(
            f"| {row['control_id']} | {_safe(row.get('synthetic_type') or row['source_type'])} | "
            f"{row['expected_label']} | {row['current_predicted_label']} | "
            f"{_probability(row['current_suspicious_probability'])} | {row['current_text_score']} | "
            f"{_probability(row['absolute_probability_delta'])} | {row['label_match']} |"
        )
    lines.extend([
        "",
        "## Panel B source/provenance",
        "",
        "- NORMAL 5: SHA-256 최소 순으로 선택한 v4.3 deduplicated human-reviewed NORMAL.",
        "- SUSPICIOUS 4: v3.2의 원래 human-reviewed SUSPICIOUS 4개와 v4.3의 exact normalized-content 교집합 전부.",
        "- Human-reviewed labels are not externally verified authenticity/fake-review facts.",
        "",
        "## Panel B results",
        "",
        "Descriptive counts only; these are not Accuracy, Recall, Precision, F1, or FPR estimates.",
        "",
        "| Panel | Human label | Length | Current label | Current p(susp.) | Text score | Label match |",
        "|---|---|---:|---|---:|---:|---|",
    ])
    for row in b["samples"]:
        lines.append(
            f"| {row['panel_id']} | {row['human_label']} | {row['character_length']} | "
            f"{row['current_predicted_label']} | {_probability(row['current_suspicious_probability'])} | "
            f"{row['current_text_score']} | {row['human_model_label_match']} |"
        )
    for label in ("NORMAL", "SUSPICIOUS"):
        summary = b["summary"][label]
        lines.append(
            f"- Human-reviewed {label} {summary['sample_count']}건: model NORMAL "
            f"{summary['model_NORMAL_count']}, model SUSPICIOUS {summary['model_SUSPICIOUS_count']}, "
            f"unavailable {summary['model_unavailable_count']}."
        )
    lines.extend([
        "",
        "## Interpretation",
        "",
        f"1. Runtime/model/config historical-label reproduction: **{a['summary']['status']}**.",
        f"2. {report['interpretation']['runtime_mismatch_assessment']}",
        f"3. {report['interpretation']['human_suspicious_observation']}",
        f"4. Recommended priority: **{report['interpretation']['recommended_priority']}**.",
        "",
        "## Limitations",
        "",
        "- Panel A validates reproduction of ten previous-model labels, not overall model quality.",
        "- Panel A positives are synthetic and are not real human-positive ground truth.",
        "- Panel B has only five NORMAL and four SUSPICIOUS examples and is not representative.",
        "- Panel B labels are human-reviewed signals, not externally verified fake/real facts.",
        "- No formal performance metric is estimated from Panel B.",
        "",
        "## Recommended next step",
        "",
        f"- {report['interpretation']['next_step_detail']}",
        "- Keep calibration, candidate expansion, and runtime investigation as separate hypotheses.",
        "",
    ])
    return "\n".join(lines)


def interpretation(panel_a: Mapping[str, object], panel_b: Mapping[str, object]) -> dict[str, str]:
    if panel_a["status"] != "PASS":
        return {
            "runtime_mismatch_assessment": "Runtime/model/config mismatch remains a priority hypothesis.",
            "human_suspicious_observation": "Panel B was not run because Panel A failed.",
            "recommended_priority": "runtime investigation",
            "next_step_detail": "Investigate runtime/model/config identity before drawing new-data conclusions.",
        }
    suspicious = panel_b["SUSPICIOUS"]
    normal_count = int(suspicious["model_NORMAL_count"])
    suspicious_count = int(suspicious["model_SUSPICIOUS_count"])
    observation = (
        f"Among four human-reviewed SUSPICIOUS examples, the model returned NORMAL {normal_count} "
        f"and SUSPICIOUS {suspicious_count}; this is descriptive, not a recall estimate."
    )
    if normal_count > suspicious_count:
        priority = "candidate expansion"
        detail = (
            "Expand a provenance-controlled human-reviewed candidate set before calibration; "
            "the four-case panel is too small for a calibration decision."
        )
    else:
        priority = "calibration investigation"
        detail = (
            "Run a larger held-out, labeled calibration study while preserving the current runtime policy."
        )
    return {
        "runtime_mismatch_assessment": (
            "Panel A provides no evidence that runtime mismatch explains the low-information audit's all-NORMAL result."
        ),
        "human_suspicious_observation": observation,
        "recommended_priority": priority,
        "next_step_detail": detail,
    }


def run(json_path: Path, markdown_path: Path) -> dict[str, object]:
    panel_a_controls = resolve_panel_a(select_panel_a(load_predictions(PREDICTIONS_PATH)), V1_SOURCE_WORKBOOK)
    panel_a_results = infer_panel_a(panel_a_controls)
    panel_a_summary = summarize_panel_a(panel_a_results)
    if panel_a_summary["status"] != "PASS":
        panel_b_controls: list[dict[str, object]] = []
        panel_b_results: list[dict[str, object]] = []
        panel_b_summary = {
            "NORMAL": {"sample_count": 0, "model_NORMAL_count": 0, "model_SUSPICIOUS_count": 0, "model_unavailable_count": 0},
            "SUSPICIOUS": {"sample_count": 0, "model_NORMAL_count": 0, "model_SUSPICIOUS_count": 0, "model_unavailable_count": 0},
        }
    else:
        panel_b_controls = select_panel_b(V3_WORKBOOK, V4_WORKBOOK)
        panel_b_results = infer_panel_b(panel_b_controls)
        panel_b_summary = summarize_panel_b(panel_b_results)
    config = json.loads((Path(DEFAULT_MODEL_PATH) / "inference_config.json").read_text(encoding="utf-8"))
    report = {
        "metadata": {
            "generated_at": datetime.now(timezone(timedelta(hours=9))).isoformat(),
            "model_path": str(Path(DEFAULT_MODEL_PATH).resolve()),
            "model_version": Path(DEFAULT_MODEL_PATH).name,
            "runtime_function": "app.analyzers.p_text.predict_text_score",
            "inference_config": config,
            "total_inference_calls": len(panel_a_results) + len(panel_b_results),
        },
        "source_resolution": {
            "panel_a_predictions": str(PREDICTIONS_PATH),
            "panel_a_content_workbook": str(V1_SOURCE_WORKBOOK),
            "panel_a_selection": "SHA-256 stable selection; exact Excel row/review ID/label validation",
            "panel_b_v3_baseline": str(V3_WORKBOOK),
            "panel_b_v4_current": str(V4_WORKBOOK),
            "panel_b_selection": (
                "v4.3 NORMAL SHA-256 stable sample of 5; all four v3.2 deduplicated "
                "SUSPICIOUS texts exact-matched into v4.3"
            ),
        },
        "panel_a": {"summary": panel_a_summary, "samples": panel_a_results},
        "panel_b": {"summary": panel_b_summary, "samples": panel_b_results},
        "interpretation": interpretation(panel_a_summary, panel_b_summary),
    }
    if report["metadata"]["total_inference_calls"] > 19:
        raise RuntimeError("inference call limit exceeded")
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    markdown_path.write_text(render_markdown(report), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--json-report", type=Path, default=DEFAULT_JSON)
    parser.add_argument("--markdown-report", type=Path, default=DEFAULT_MARKDOWN)
    args = parser.parse_args()
    import torch
    torch.set_num_threads(1)
    report = run(args.json_report, args.markdown_report)
    print(json.dumps({
        "panel_a": report["panel_a"]["summary"],
        "panel_b": report["panel_b"]["summary"],
        "total_inference_calls": report["metadata"]["total_inference_calls"],
        "recommended_priority": report["interpretation"]["recommended_priority"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
